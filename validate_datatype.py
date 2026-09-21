import os
import logging
from datetime import datetime
import openpyxl
import config
from utility import Utility

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
log_filename = f"datatype_log_{timestamp}.txt"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    handlers=[
        logging.FileHandler(log_filename, encoding="utf-8"),
        logging.StreamHandler()
    ]
)

util = Utility()

def get_datatype_set(datatype_string):
    """Converts a comma-separated string into a set of trimmed datatypes."""
    if not datatype_string:
        return set()
    return set(item.strip() for item in str(datatype_string).split(',') if item.strip())

def validate_standard_folder(filepath, folder_name, expected_rule):
    """Validates files that do not require multi-file aggregation."""
    filename = os.path.basename(filepath)
    
    try:
        workbook = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        if 'REQUEST_TABLE' not in workbook.sheetnames:
            return f"[ERROR] Sheet 'REQUEST_TABLE' missing in {filename}"
        
        worksheet = workbook['REQUEST_TABLE']

        if folder_name == "annual_financials_insurance":
            expected_set = get_datatype_set(expected_rule)
            current_row = 7
            has_data = False
            
            while worksheet[f"G{current_row}"].value is not None:
                has_data = True
                cell_value = worksheet[f"F{current_row}"].value
                actual_set = get_datatype_set(cell_value)
                
                if actual_set != expected_set:
                    missing = expected_set - actual_set
                    return f"[MISMATCH] {filename} row {current_row}: Missing {len(missing)} datatypes: {', '.join(missing)}"
                
                current_row += 1
                
            if not has_data:
                return f"[ERROR] Request table starting at row 7 is empty for {filename}"
                
        else:
            # For daily/quarterly dictionaries (row index -> expected string)
            for row_index, expected_string in expected_rule.items():
                expected_set = get_datatype_set(expected_string)
                cell_value = worksheet[f"F{row_index}"].value
                actual_set = get_datatype_set(cell_value)
                
                if actual_set != expected_set:
                    missing = expected_set - actual_set
                    return f"[MISMATCH] {filename} row {row_index}: Missing {len(missing)} datatypes: {', '.join(missing)}"
                    
        return f"[OK] {filename}"

    except Exception as exception:
        return f"[ERROR] Failed to read {filename}: {str(exception)}"
    finally:
        if 'workbook' in locals():
            workbook.close()

base_directory = '.' 
error_list = []

logging.info("Starting datatype validation process...")

for root_directory, directories, files in os.walk(base_directory):
    directories[:] = [directory for directory in directories if directory not in config.EXCLUDED_FOLDERS]
    folder_name = os.path.basename(root_directory).lower()
    
    expected_rule = config.EXPECTED_DATATYPES.get(folder_name)
    if not expected_rule:
        continue

    logging.info(f"===== Scanning directory: {root_directory} =====")

    if folder_name == "annual_financials_global":
        # Group files by country first
        country_groups = {}
        for file_name in files:
            if file_name.endswith('.xlsm') and not file_name.startswith('~'):
                file_info = util.parse_filename(file_name)
                if file_info:
                    country = file_info["country"]
                    if country not in country_groups:
                        country_groups[country] = []
                    country_groups[country].append(file_name)

        expected_global_set = get_datatype_set(expected_rule)
        
        for country, country_files in country_groups.items():
            aggregated_years = {}
            
            # Aggregate all files for THIS country
            for file_name in country_files:
                filepath = os.path.join(root_directory, file_name)
                file_info = util.parse_filename(file_name)
                start_year = file_info["start_year"]
                end_year = file_info["end_year"] if file_info["end_year"] else start_year
                
                try:
                    workbook = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
                    if 'REQUEST_TABLE' in workbook.sheetnames:
                        worksheet = workbook['REQUEST_TABLE']
                        file_years_range = list(range(start_year, end_year + 1))
                        
                        for index, file_year in enumerate(file_years_range):
                            cell_value = worksheet[f"F{7 + index}"].value
                            actual_set = get_datatype_set(cell_value)
                            
                            if file_year not in aggregated_years:
                                aggregated_years[file_year] = set()
                            aggregated_years[file_year].update(actual_set)
                            
                        logging.info(f"[SCANNED] {file_name} successfully aggregated.")
                    else:
                        error_msg = f"[ERROR] Sheet 'REQUEST_TABLE' missing in {file_name}"
                        logging.error(error_msg)
                        error_list.append(error_msg)
                except Exception as exception:
                    error_msg = f"[ERROR] Failed to read {file_name}: {str(exception)}"
                    logging.error(error_msg)
                    error_list.append(error_msg)
                finally:
                    if 'workbook' in locals():
                        workbook.close()
            
            # Immediately validate THIS country right after its files are aggregated
            for year, actual_set in aggregated_years.items():
                if actual_set != expected_global_set:
                    missing = expected_global_set - actual_set
                    error_message = f"[MISMATCH] annual_financials_global - {country} ({year}): Missing {len(missing)} required datatypes: {', '.join(missing)}"
                    logging.error(error_message)
                    error_list.append(error_message)
                else:
                    logging.info(f"[OK] annual_financials_global - {country} ({year}) has complete datatypes.")

    else:
        # Standard file-by-file check for all other folders
        for file_name in files:
            if file_name.endswith('.xlsm') and not file_name.startswith('~'):
                filepath = os.path.join(root_directory, file_name)
                result = validate_standard_folder(filepath, folder_name, expected_rule)
                if result.startswith('[MISMATCH]') or result.startswith('[ERROR]'):
                    logging.error(result)
                    error_list.append(result)
                elif result.startswith('[SKIPPED]'):
                    logging.warning(result)
                else:
                    logging.info(result)

if error_list:
    logging.warning("--- Errors Found ---")
    with open('datatype_manual_error_list.txt', 'w', encoding='utf-8') as error_file:
        for error_message in error_list:
            error_file.write(error_message + '\n')
    logging.info("Errors saved to datatype_manual_error_list.txt")
else:
    logging.info("\n--- All checks passed ---")
